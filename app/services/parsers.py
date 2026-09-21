from pathlib import Path
from typing import Any
import io

from docx import Document
from pypdf import PdfReader
from pptx import Presentation
from PIL import Image
import pytesseract
import fitz  # PyMuPDF


TEXT_EXTS = {".txt", ".md", ".csv"}


def extract_text(path: str | Path) -> str:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".docx":
        doc = Document(path)
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    if suffix == ".pdf":
        reader = PdfReader(str(path))
        text = "\n".join((p.extract_text() or "") for p in reader.pages)
        if text.strip():
            return text
        return ocr_pdf(path)
    if suffix in {".pptx", ".ppt"}:
        return extract_pptx_text(path)
    if suffix in {".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".webp"}:
        return pytesseract.image_to_string(Image.open(path))
    if suffix in TEXT_EXTS:
        return path.read_text(encoding="utf-8", errors="ignore")
    return ""


def extract_pptx_text(path: str | Path) -> str:
    prs = Presentation(str(path))
    out = []
    for i, slide in enumerate(prs.slides, start=1):
        out.append(f"Slide {i}")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                txt = shape.text.strip()
                if txt:
                    out.append(txt)
    return "\n".join(out)


def ocr_pdf(path: str | Path) -> str:
    doc = fitz.open(str(path))
    pages = []
    for page in doc:
        pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5))
        image = Image.open(io.BytesIO(pix.tobytes("png")))
        pages.append(pytesseract.image_to_string(image))
    doc.close()
    return "\n".join(pages)


def analyze_docx(path: str | Path) -> dict[str, Any]:
    doc = Document(str(path))
    fonts = []
    for paragraph in doc.paragraphs[:50]:
        for run in paragraph.runs:
            if run.font.name:
                fonts.append(run.font.name)
    section = doc.sections[0] if doc.sections else None
    return {
        "type": "docx",
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "fonts": list(dict.fromkeys(fonts))[:5],
        "styles": list(dict.fromkeys(p.style.name for p in doc.paragraphs if p.style))[:10],
        "page_size_inches": {
            "width": round(section.page_width.inches, 2) if section else None,
            "height": round(section.page_height.inches, 2) if section else None,
        },
        "text_preview": extract_text(path)[:1500],
    }


def analyze_pptx(path: str | Path) -> dict[str, Any]:
    prs = Presentation(str(path))
    fonts = []
    layouts = []
    titles = []
    for slide in list(prs.slides)[:10]:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                if shape.text.strip():
                    titles.append(shape.text.strip().splitlines()[0][:120])
                for p in shape.text_frame.paragraphs:
                    for run in p.runs:
                        if run.font.name:
                            fonts.append(run.font.name)
        if slide.slide_layout:
            layouts.append(slide.slide_layout.name)
    return {
        "type": "pptx",
        "slide_count": len(prs.slides),
        "slide_size_inches": {
            "width": round(prs.slide_width.inches, 2),
            "height": round(prs.slide_height.inches, 2),
        },
        "layouts_seen": list(dict.fromkeys(layouts))[:10],
        "fonts": list(dict.fromkeys(fonts))[:5],
        "text_preview": extract_text(path)[:1500],
        "shape_text_preview": titles[:10],
    }


def analyze_file(path: str | Path) -> dict[str, Any]:
    suffix = Path(path).suffix.lower()
    if suffix == ".docx":
        return analyze_docx(path)
    if suffix in {".pptx", ".ppt"}:
        return analyze_pptx(path)
    return {
        "type": suffix.lstrip("."),
        "text_preview": extract_text(path)[:1500],
    }
