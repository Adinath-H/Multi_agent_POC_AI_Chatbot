from pathlib import Path
import re

from docx import Document
from docx.shared import Pt
from pptx import Presentation
from pptx.util import Inches, Pt as PPTPt

from .parsers import extract_text


def clean_filename(text: str) -> str:
    text = re.sub(r"[^A-Za-z0-9._-]+", "_", text.strip())
    return text[:80] or "artifact"


def generate_docx(content: str, output_path: str, template_path: str | None = None):
    document = Document(template_path) if template_path else Document()
    if document.paragraphs and document.paragraphs[-1].text.strip():
        document.add_page_break()

    for block in re.split(r"\n\s*\n", content.strip()):
        lines = [x.strip() for x in block.splitlines() if x.strip()]
        if not lines:
            continue
        first = re.sub(r"^#+\s*", "", lines[0])
        if lines[0].startswith("# "):
            p = document.add_heading(first, level=0)
        elif lines[0].startswith("## "):
            p = document.add_heading(first, level=1)
        elif lines[0].startswith("### "):
            p = document.add_heading(first, level=2)
        elif len(lines) == 1:
            p = document.add_paragraph(first)
        else:
            p = document.add_paragraph()
            for i, line in enumerate(lines):
                is_bullet = bool(re.match(r"^[-*•]\s+", line))
                clean = re.sub(r"^[-*•]\s*", "", line)
                run = p.add_run(("• " if is_bullet else "") + clean)
                if i < len(lines) - 1:
                    run.add_break()
    document.save(output_path)
    return output_path


def parse_slides(content: str) -> list[tuple[str, list[str]]]:
    slides = []
    parts = re.split(r"\[SLIDE\s+\d+\]", content, flags=re.I)
    for part in parts:
        if not part.strip():
            continue
        lines = [x.strip() for x in part.strip().splitlines() if x.strip()]
        if not lines:
            continue
        title = lines[0].lstrip("# ")
        bullets = []
        for line in lines[1:]:
            line = re.sub(r"^[-*•]\s*", "", line)
            if line:
                bullets.append(line)
        slides.append((title, bullets[:6]))
    return slides


def generate_pptx(content: str, output_path: str, template_path: str | None = None):
    prs = Presentation(template_path) if template_path else Presentation()
    if template_path and len(prs.slides) > 0:
        layout = prs.slide_layouts[1] if len(prs.slide_layouts) > 1 else prs.slide_layouts[0]
    else:
        layout = prs.slide_layouts[1]

    for title, bullets in parse_slides(content):
        slide = prs.slides.add_slide(layout)
        if slide.shapes.title:
            slide.shapes.title.text = title
        else:
            box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11), Inches(0.8))
            box.text_frame.paragraphs[0].text = title
        body = None
        for ph in slide.placeholders:
            try:
                if ph.placeholder_format.idx != 0 and ph.has_text_frame:
                    body = ph
                    break
            except Exception:
                pass
        if body:
            tf = body.text_frame
            tf.clear()
            for i, bullet in enumerate(bullets or ["Key point"]):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = bullet
                p.level = 0
                for run in p.runs:
                    run.font.size = PPTPt(24)
    prs.save(output_path)
    return output_path


def docx_to_pptx(docx_path: str, output_path: str):
    text = extract_text(docx_path)
    sections = re.split(r"\n(?=#{1,3}\s)", text)
    chunks = []
    for s in sections:
        lines = [x.strip() for x in s.splitlines() if x.strip()]
        if not lines:
            continue
        title = lines[0].lstrip("# ")[:100]
        bullets = [re.sub(r"^[-*]\s*", "", x) for x in lines[1:6]]
        chunks.append((title, bullets))
    if not chunks:
        chunks = [("Converted Document", [text[:500]])]
    content = "\n\n".join(
        [f"[SLIDE {i}]\n{title}\n" + "\n".join(f"- {b}" for b in bullets)
         for i, (title, bullets) in enumerate(chunks[:15], start=1)]
    )
    return generate_pptx(content, output_path)


def pptx_to_docx(pptx_path: str, output_path: str):
    prs = Presentation(pptx_path)
    doc = Document()
    for i, slide in enumerate(prs.slides, start=1):
        title = "Slide " + str(i)
        slide_text = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                txt = shape.text.strip()
                if txt:
                    slide_text.append(txt)
        if slide_text:
            title = slide_text[0][:100]
        doc.add_heading(title, level=1)
        for text in slide_text[1:]:
            doc.add_paragraph(text)
    doc.save(output_path)
    return output_path
